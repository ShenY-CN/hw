# 本程序及代码在人工智能工具辅助下完成：OpenAI Codex（GPT-6，OpenAI；GPT-6 模型家族发布日期 2026-09-03）。
# 参赛队须自行理解、复核与改写；本会话未提供更细的子型号标识。
"""Run the entire calculation in a clean copy and record its provenance."""
import argparse
import hashlib
import importlib.metadata
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
from datetime import datetime,timezone

ROOT=Path(__file__).resolve().parents[1]
STEPS=['common.py','solve_transport.py','solve_joint.py',
       'partition_validate.py','energy_sensitivity.py',
       'export_submission.py','figures.py','render_roadmap.py']
PACKAGES=['numpy','openpyxl','scipy','rasterio','pyproj','ortools','matplotlib','affine']


def digest(path):
    h=hashlib.sha256()
    with Path(path).open('rb') as stream:
        for chunk in iter(lambda:stream.read(1024*1024),b''):
            h.update(chunk)
    return h.hexdigest()


def hashes(folder):
    return {str(p.relative_to(folder)):digest(p) for p in sorted(folder.rglob('*')) if p.is_file() and p.name!='.DS_Store'}


def run(out):
    out=Path(out).resolve()
    if out.exists():raise FileExistsError(f'Choose an empty output directory: {out}')
    locked={}
    for line in (ROOT/'requirements.lock').read_text(encoding='utf8').splitlines():
        if line.strip():
            name,version=line.split('==',1)
            locked[name]=version
    installed={name:importlib.metadata.version(name) for name in PACKAGES}
    mismatches={name:(installed[name],version) for name,version in locked.items()
                if installed.get(name)!=version}
    if mismatches:raise RuntimeError(f'Dependency versions differ from requirements.lock: {mismatches}')
    out.mkdir(parents=True)
    shutil.copytree(ROOT/'code',out/'code',ignore=shutil.ignore_patterns('__pycache__','*.pyc','.DS_Store'))
    shutil.copytree(ROOT/'input',out/'input',ignore=shutil.ignore_patterns('.DS_Store'))
    shutil.copy2(ROOT/'requirements.lock',out/'requirements.lock')
    (out/'results').mkdir();(out/'figures').mkdir()
    shutil.copy2(ROOT/'figures'/'fig_roadmap.drawio',out/'figures'/'fig_roadmap.drawio')
    env=dict(os.environ,PYTHONDONTWRITEBYTECODE='1',MPLCONFIGDIR=str(out/'matplotlib_cache'))
    logs=[]
    for step in STEPS:
        command=[sys.executable,str(out/'code'/step)]
        print('RUN',step,flush=True)
        completed=subprocess.run(command,cwd=out,env=env,text=True,stdout=subprocess.PIPE,
                                 stderr=subprocess.STDOUT)
        logs.append(dict(command=command,exit_code=completed.returncode,output=completed.stdout))
        (out/'run.log').write_text('\n\n'.join(
            f"{x['command']}\nexit={x['exit_code']}\n{x['output']}" for x in logs),encoding='utf8')
        if completed.returncode:
            print(completed.stdout[-3000:],flush=True)
            raise RuntimeError(f'{step} failed with exit {completed.returncode}; see {out/"run.log"}')
    validation=json.loads((out/'results'/'validation.json').read_text(encoding='utf8'))
    if not all(validation[str(q)]['pass_'] for q in (2,3)):
        raise RuntimeError(f'Constraint check failed: {validation}')
    q2=json.loads((out/'results'/'q2.json').read_text(encoding='utf8'))['summary']
    q3=json.loads((out/'results'/'q3.json').read_text(encoding='utf8'))['summary']
    q4=json.loads((out/'results'/'q4.json').read_text(encoding='utf8'))
    from openpyxl import load_workbook
    book=load_workbook(out/'D题结果提交表.xlsx',read_only=True,data_only=True)
    if len(book.sheetnames)!=8:raise RuntimeError('Submission workbook lacks required sheets')
    pdfs=sorted(p.stem for p in (out/'figures').glob('*.pdf'))
    pngs=sorted(p.stem for p in (out/'figures').glob('*.png'))
    if len(pdfs)!=9 or sorted(name for name in pdfs if name!='fig_roadmap')!=pngs:
        raise RuntimeError(f'Figure outputs incomplete: PDF={pdfs}, PNG={pngs}')
    manifest=dict(
        created_at=datetime.now(timezone.utc).isoformat(),python=sys.version.split()[0],
        packages=installed,requirements_sha256=digest(out/'requirements.lock'),
        random_seeds=dict(route_construction=list(range(24)),cp_sat=42),
        commands=[dict(script=s,exit_code=0) for s in STEPS],
        input_sha256=hashes(out/'input'),code_sha256=hashes(out/'code'),
        roadmap_source_sha256=digest(out/'figures'/'fig_roadmap.drawio'),
        q2=q2,q3=q3,q4={k:v['selected']['total'] for k,v in q4['schemes'].items()},
        workbook_sha256=digest(out/'D题结果提交表.xlsx'),
        figure_files=[name+'.pdf' for name in pdfs])
    (out/'results'/'复现清单.json').write_text(
        json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf8')
    print('PASS',out,flush=True)
    print('Q2',q2,flush=True)
    print('Q3',q3,flush=True)
    return out


if __name__=='__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('--output',required=True,help='New directory for an isolated fresh run')
    args=parser.parse_args()
    run(args.output)
