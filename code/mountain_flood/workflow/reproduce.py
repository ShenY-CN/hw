# 本程序及代码在人工智能工具辅助下完成：OpenAI Codex（GPT-5，OpenAI；GPT-5 发布于 2025-08-07）。
# 参赛队须自行理解、复核与改写；本会话未提供更细的子型号标识。
"""重算问题一，并独立回放和验证归档的搜索候选方案。"""

# 支持直接运行本文件；此处将 code/ 加入模块搜索路径。
if __package__ in (None, ""):
    import sys
    from pathlib import Path as _BootstrapPath
    sys.path.insert(0, str(_BootstrapPath(__file__).resolve().parents[2]))
import argparse
import hashlib
import importlib.metadata
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
from datetime import datetime,timezone

from mountain_flood.paths import ROOT
STEPS=[
    'mountain_flood/core/domain.py',
    'mountain_flood/problem1/recompute.py',
    'mountain_flood/validation/verify_methods_archive.py',
    'mountain_flood/experiments/formal_q2_protocol.py',
    'mountain_flood/experiments/formal_q3_protocol.py',
    'mountain_flood/validation/compare_joint_candidates.py',
    'mountain_flood/workflow/promote_fusion.py',
    'mountain_flood/validation/runner.py',
    'independent_solution_audit.py',
    'audit_kmax_archive.py',
    'controlled_stress.py',
    'mountain_flood/experiments/energy_sensitivity.py',
    'mountain_flood/reporting/export_submission.py',
    'mountain_flood/figure/q1.py',
    'mountain_flood/figure/q2.py',
    'mountain_flood/figure/q3.py',
    'mountain_flood/figure/q4.py',
    'mountain_flood/figure/spatial.py',
    'mountain_flood/figure/overview.py',
    'figure_raw_inputs.py',
    'assemble_figure_pool.py',
]
PACKAGES=['numpy','openpyxl','scipy','rasterio','pyproj','ortools','matplotlib','affine']


def digest(path):
    h=hashlib.sha256()
    with Path(path).open('rb') as stream:
        for chunk in iter(lambda:stream.read(1024*1024),b''):
            h.update(chunk)
    return h.hexdigest()


def hashes(folder):
    return {str(p.relative_to(folder)):digest(p) for p in sorted(folder.rglob('*')) if p.is_file() and p.name!='.DS_Store'}


def run(out):
    out=Path(out).resolve()
    if out.exists():raise FileExistsError(f'Choose an empty output directory: {out}')
    locked={}
    for line in (ROOT/'requirements.lock').read_text(encoding='utf8').splitlines():
        if line.strip():
            name,version=line.split('==',1)
            locked[name]=version
    installed={name:importlib.metadata.version(name) for name in PACKAGES}
    mismatches={name:(installed[name],version) for name,version in locked.items()
                if installed.get(name)!=version}
    if mismatches:raise RuntimeError(f'Dependency versions differ from requirements.lock: {mismatches}')
    out.mkdir(parents=True)
    shutil.copytree(ROOT/'code',out/'code',ignore=shutil.ignore_patterns('__pycache__','*.pyc','.DS_Store'))
    shutil.copytree(ROOT/'input',out/'input',ignore=shutil.ignore_patterns('.DS_Store','~$*'))
    shutil.copytree(ROOT/'config',out/'config')
    shutil.copy2(ROOT/'requirements.lock',out/'requirements.lock')
    (out/'results').mkdir();(out/'figures').mkdir()
    shutil.copy2(ROOT/'figures'/'fig_roadmap.drawio',out/'figures'/'fig_roadmap.drawio')
    archives=['q2_time_energy_candidate.json','q3_hard_hill1_e3900_margin.json',
              'method_comparison.json','method_comparison_strict_archive.json',
              'q3_soft_legacy.json','kmax_scan_20260926.json']
    for pattern in ('q3_hard_hill1_e*_trim.json','q3_hard_newq2_e*_trim.json',
                    'q3_hard_q2_e*_completion_trim.json','q3_hard_q2_e*_buffer15_trim.json',
                    'q3_hard_hill1_e*_margin.json'):
        archives.extend(p.name for p in sorted((ROOT/'results').glob(pattern)))
    archives=list(dict.fromkeys(archives))
    for filename in archives:
        shutil.copy2(ROOT/'results'/filename,out/'results'/filename)
    figure_set='reproduced_figures'
    env=dict(os.environ,PYTHONDONTWRITEBYTECODE='1',MPLCONFIGDIR=str(out/'matplotlib_cache'),
             MOUNTAIN_FLOOD_FIGURE_SET=figure_set)
    logs=[]
    for step in STEPS:
        command=[sys.executable,str(out/'code'/step)]
        print('RUN',step,flush=True)
        completed=subprocess.run(command,cwd=out,env=env,text=True,stdout=subprocess.PIPE,
                                 stderr=subprocess.STDOUT)
        logs.append(dict(command=command,exit_code=completed.returncode,output=completed.stdout))
        (out/'run.log').write_text('\n\n'.join(
            f"{x['command']}\nexit={x['exit_code']}\n{x['output']}" for x in logs),encoding='utf8')
        if completed.returncode:
            print(completed.stdout[-3000:],flush=True)
            raise RuntimeError(f'{step} failed with exit {completed.returncode}; see {out/"run.log"}')
    validation=json.loads((out/'results'/'validation.json').read_text(encoding='utf8'))
    if not all(validation[str(q)]['pass_'] for q in (2,3)):
        raise RuntimeError(f'Constraint check failed: {validation}')
    q2=json.loads((out/'results'/'q2.json').read_text(encoding='utf8'))['summary']
    q3=json.loads((out/'results'/'q3.json').read_text(encoding='utf8'))['summary']
    q4=json.loads((out/'results'/'q4.json').read_text(encoding='utf8'))
    from openpyxl import load_workbook
    book=load_workbook(out/'D题结果提交表.xlsx',read_only=True,data_only=True)
    if len(book.sheetnames)!=8:raise RuntimeError('Submission workbook lacks required sheets')
    figure_folder=out/'figures'/figure_set
    pdfs=sorted(p.stem for p in figure_folder.glob('*.pdf'))
    pngs=sorted(p.stem for p in figure_folder.glob('*.png'))
    svgs=sorted(p.stem for p in figure_folder.glob('*.svg'))
    expected_figures={
        'q1_dem_nodes','q1_safe_capacity','q1_capacity_heatmap','q1_typical_elevation_profile',
        'q1_payload_energy','q1_objective_solutions','q1_sensitivity','q1_return_soc',
        'q2_transport_routes','q2_drone_timeline',
        'q2_battery_timeline','q2_delivery_deadlines','q2_return_soc',
        'q2_search_comparison','q2_route_energy','q3_relay_map','q3_relay_coverage_points',
        'q3_communication_timeline','q3_link_margin','q3_joint_timeline','q3_battery_timeline',
        'q3_relay_soc','q3_q2_q3_comparison','q4_partition_2groups',
        'q4_partition_3groups','q4_resource_demand','q4_workload',
        'q4_resource_deficit','q4_task_network',
        'spatial_terrain_time_partitions','spatial_group_standard_ellipses',
        'spatial_terrain_routes_3d','spatial_time_sliced_routes',
    }
    expected_figures.add('fig_roadmap')
    if set(pdfs)!=expected_figures or set(pngs)!=expected_figures or set(svgs)!=expected_figures:
        raise RuntimeError(f'Figure outputs incomplete: PDF={pdfs}, PNG={pngs}, SVG={svgs}')
    pool=out/'figures'/'rigor_pool_20260926'
    pool_manifest=json.loads((pool/'figure_pool_manifest.json').read_text(encoding='utf8'))
    if len(pool_manifest['file_sha256'])!=39 or any(
            digest(pool/name)!=recorded for name,recorded in pool_manifest['file_sha256'].items()):
        raise RuntimeError('Audited figure pool is incomplete or its manifest differs')
    kmax_audit=json.loads((out/'results'/'kmax_archive_audit_20260926.json').read_text(encoding='utf8'))
    if kmax_audit['status']!='PASS' or kmax_audit['passed']!=48:
        raise RuntimeError('Archived K-scan route witnesses failed independent audit')
    verified_results=('q2_formal_protocol.json','q3_formal_coordination_audit.json','q4.json')
    manifest=dict(
        created_at=datetime.now(timezone.utc).isoformat(),python=sys.version.split()[0],
        packages=installed,requirements_sha256=digest(out/'requirements.lock'),
        random_seeds=dict(route_construction=list(range(24)),method_comparison=[0,1,2],cp_sat=42),
        archived_selected_plans_sha256={name:digest(out/'results'/name) for name in archives},
        verified_results_sha256={name:digest(out/'results'/name) for name in verified_results},
        archive_note='Q1 is recalculated; revised-rule and K-limit wall-clock searches, prior strict-rule search, historical soft plan, and finite joint search are archived as route witnesses. All 48 K-limit witnesses and current plans are independently recalculated and validated. This replay does not repeat wall-clock discovery.',
        commands=[dict(script=s,exit_code=0) for s in STEPS],
        input_sha256=hashes(out/'input'),code_sha256=hashes(out/'code'),config_sha256=hashes(out/'config'),
        roadmap_source_sha256=digest(out/'figures'/'fig_roadmap.drawio'),
        q2=q2,q3=q3,q4={k:v['selected']['total'] for k,v in q4['schemes'].items()},
        workbook_sha256=digest(out/'D题结果提交表.xlsx'),
        figure_files=[f'{figure_set}/{name}.pdf' for name in pdfs],
        audited_figure_pool_sha256=pool_manifest['file_sha256'],
        kmax_archive_audit_sha256=digest(out/'results'/'kmax_archive_audit_20260926.json'))
    (out/'results'/'复现清单.json').write_text(
        json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf8')
    print('PASS',out,flush=True)
    print('Q2',q2,flush=True)
    print('Q3',q3,flush=True)
    return out


if __name__=='__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('--output',required=True,help='New directory for an isolated fresh run')
    parser.add_argument('--manifest-target',help='Optional path receiving the generated manifest')
    args=parser.parse_args()
    completed=run(args.output)
    if args.manifest_target:
        shutil.copy2(completed/'results'/'复现清单.json',Path(args.manifest_target).resolve())
