# 本程序及代码在人工智能工具辅助下完成：OpenAI Codex（GPT-5，OpenAI；GPT-5 发布于 2025-08-07）。
# 参赛队须自行理解、复核与改写；本会话未提供更细的子型号标识。
"""从同一组模型结果 JSON 填写提交模板，并回读核对导出内容。"""

# 支持直接运行本文件；此处将 code/ 加入模块搜索路径。
if __package__ in (None, ""):
    import sys
    from pathlib import Path as _BootstrapPath
    sys.path.insert(0, str(_BootstrapPath(__file__).resolve().parents[2]))
from pathlib import Path
import argparse
import json
from openpyxl import load_workbook

from mountain_flood.paths import ROOT, RESULT


def load(name):
    return json.loads((RESULT/name).read_text(encoding='utf8'))


def put(sheet,records):
    for row_index,record in enumerate(records,2):
        for col_index,value in enumerate(record,1):
            sheet.cell(row_index,col_index,value)


def check_rows(sheet,expected):
    actual=[row[:len(expected[0])] for row in list(sheet.values)[1:] if row[0] is not None]
    if len(actual)!=len(expected):
        raise RuntimeError(f'{sheet.title}: {len(actual)} rows, expected {len(expected)}')
    for i,(row,reference) in enumerate(zip(actual,expected),2):
        for col,(value,wanted) in enumerate(zip(row,reference),1):
            if isinstance(wanted,(int,float)) and not isinstance(wanted,bool):
                good=isinstance(value,(int,float)) and abs(value-wanted)<=1e-6
            else:
                good=value==wanted
            if not good:raise RuntimeError(f'{sheet.title}!{i}:{col}: {value!r} != {wanted!r}')


def trip_rows(routes):
    return [(r['id'],r['unit'],r['g'],r['battery'],r['start'],
             '→'.join(f'S{n:03}' for n in r['order']),r['end'],r['energy'])
            for r in routes]


def delivery_rows(routes,boxes):
    by_box={}
    for r in routes:
        for index,time in r['deliver'].items():
            i=int(index)
            if i in by_box:raise ValueError(f'Duplicate box {i}')
            by_box[i]=(r['id'],r['start']+time)
    if set(by_box)!=set(range(len(boxes))):raise ValueError('Missing box delivery')
    return [(b['id'],by_box[i][0],f"S{b['node']:03}",by_box[i][1])
            for i,b in enumerate(boxes)]


def communication_rows(records):
    merged=[]
    for c in records:
        key=(c['route'],c['phase'],c['provider'],c['relay_task'])
        if merged and merged[-1][0]==key and abs(merged[-1][2]-c['start'])<1e-6:
            merged[-1]=(key,merged[-1][1],c['end'])
        else:
            merged.append((key,c['start'],c['end']))
    return [(key[0],key[1],start,end,
             '固定网关直连' if key[2]=='G01' else '中继保障',key[3])
            for key,start,end in merged]


def build(output):
    boxes=load('data_profile.json')['boxes']
    q1=load('q1_rho20.json');q2=load('q2.json')
    q3=load('q3.json');q4=load('q4.json')
    template=ROOT/'input'/'结果提交模板.xlsx'
    workbook=load_workbook(template)
    expected=['Q1_单点组批','Q2_运输架次','Q2_逐箱交付',
              'Q3_中继架次','Q3_通信保障','Q4_分区配置']
    if workbook.sheetnames!=expected:raise ValueError(f'Template sheets changed: {workbook.sheetnames}')
    q1_rows=[
        (f'Q1-{i:03}',f"S{r['order'][0]:03}",r['g'],
         '、'.join(boxes[int(b)]['id'] for b in r['boxes']),r['w'],r['v'],
         r['duration'],r['energy'],100*r['soc'])
        for i,r in enumerate(q1['routes'],1)]
    q2_trip_rows=trip_rows(q2['routes'])
    q2_delivery_rows=delivery_rows(q2['routes'],boxes)
    relay_rows=[
        (r['id'],r['unit'],r['component'],r['start'],*r['pos'],
         r['ready'],r['service_end'],r['end'],r['energy'])
        for r in q3['relays']]
    comm_rows=communication_rows(q3['communication'])
    group_rows=[]
    for k in (2,3):
        groups=q4['schemes'][str(k)]['selected']['groups']
        for i,g in enumerate(groups,1):
            group_rows.append((k,f'G{i}','、'.join(f'S{n:03}' for n in g['nodes']),*g['resources']))
    rows_by_sheet=dict(zip(expected,[q1_rows,q2_trip_rows,q2_delivery_rows,
                                      relay_rows,comm_rows,group_rows]))
    for name,records in rows_by_sheet.items():put(workbook[name],records)
    transport=workbook.copy_worksheet(workbook[expected[1]])
    transport.title='Q3_运输架次'
    q3_trip_rows=trip_rows(q3['routes'])
    put(transport,q3_trip_rows)
    deliveries=workbook.copy_worksheet(workbook[expected[2]])
    deliveries.title='Q3_逐箱交付'
    q3_delivery_rows=delivery_rows(q3['routes'],boxes)
    put(deliveries,q3_delivery_rows)
    rows_by_sheet.update({'Q3_运输架次':q3_trip_rows,'Q3_逐箱交付':q3_delivery_rows})
    output=Path(output)
    output.parent.mkdir(parents=True,exist_ok=True)
    workbook.save(output)
    check=load_workbook(output,read_only=True,data_only=True)
    if len(check.sheetnames)!=8:raise RuntimeError('Result workbook missing sheets')
    for name,records in rows_by_sheet.items():check_rows(check[name],records)
    return output


if __name__=='__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('--output',default=str(ROOT/'D题结果提交表.xlsx'))
    args=parser.parse_args()
    print(build(args.output))
